"""Makes sample Word templates with {{ fields }} for render_docx().
DEMO TEMPLATES ONLY. Not legal documents. A firm would supply its own.
Run: uv run python scripts/make_templates.py
"""
from pathlib import Path

from docx import Document
from docx.shared import Pt

OUT = Path(__file__).resolve().parent.parent / "data" / "templates"
OUT.mkdir(parents=True, exist_ok=True)


def doc_with(title: str, paragraphs: list[str], path: Path) -> None:
    d = Document()
    d.styles["Normal"].font.name = "Calibri"
    d.styles["Normal"].font.size = Pt(11)
    d.add_heading(title, level=1)
    for p in paragraphs:
        d.add_paragraph(p)  # one run per paragraph keeps {{ fields }} intact
    d.save(path)


doc_with("Contingency Fee Agreement (SAMPLE - not for real use)", [
    "{{ firm_name }}",
    "Date: {{ date }}",
    "Client: {{ client_name }}, {{ client_address }}",
    "Matter: Injuries from the incident on {{ accident_date }}: {{ accident_description }}",
    "1. The Client hires the Firm to pursue claims arising from the incident above.",
    "2. Fee: {{ fee_percent_pre }}% of any recovery if resolved before a lawsuit is filed, "
    "or {{ fee_percent_lit }}% after a lawsuit is filed. No recovery, no fee.",
    "3. Costs advanced by the Firm are repaid from any recovery.",
    "Client signature: ______________________   Date: __________",
    "Attorney signature: ____________________   Date: __________",
], OUT / "retainer_agreement.docx")

doc_with("Letter of Representation (SAMPLE)", [
    "{{ date }}",
    "{{ insurer_name }}",
    "Re: Our client {{ client_name }}; Date of loss: {{ accident_date }}; Claim No.: {{ claim_number }}",
    "Please be advised that this firm represents {{ client_name }} for injuries from the incident on "
    "{{ accident_date }}. Direct all communication to our office and do not contact our client.",
    "Please confirm the policy limits for your insured and preserve all evidence related to this incident.",
    "Sincerely,",
    "{{ attorney_name }}, {{ firm_name }}",
], OUT / "letter_of_representation.docx")

print(f"Made templates in {OUT}")
